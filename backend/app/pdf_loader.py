import os
import pickle
from typing import List

from docx import Document
from pptx import Presentation

from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document as LangChainDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_community.retrievers import BM25Retriever

from app.config import VECTOR_STORE_DIR


embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-l6-v2")

BM25_FILE_PATH = os.path.join(VECTOR_STORE_DIR, "bm25_store.pkl")

def load_docx(file_path: str) -> List[LangChainDocument]:

    doc = Document(file_path)

    full_text = [para.text
        for para in doc.paragraphs
        if para.text.strip()]
    text = "\n".join(full_text)

    return [LangChainDocument(page_content=text, metadata={"source": os.path.basename(file_path)})]


def load_pptx(file_path: str) -> List[LangChainDocument]:

    prs = Presentation(file_path)

    text_runs = []

    for slide_idx, slide in enumerate(prs.slides):

        for shape in slide.shapes:
            if (hasattr(shape, "text") and shape.text.strip()):
                text_runs.append(f"[Slide {slide_idx + 1}] " f"{shape.text}")

    text = "\n".join(text_runs)

    return [LangChainDocument(page_content=text,metadata={"source": os.path.basename(file_path)})]

def load_txt(file_path: str) -> List[LangChainDocument]:
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()
        return [LangChainDocument(page_content=text, metadata={"source": os.path.basename(file_path)})]



def process_pdfs_and_create_retrievers(file_paths: List[str]):
    all_docs = []

    for path in file_paths:
        ext = os.path.splitext(path)[1].lower()
        file_name = os.path.basename(path)

        if ext == ".pdf":
            loader = PyPDFLoader(path)
            docs = loader.load()

            for doc in docs:
                doc.metadata["source"] = file_name

            all_docs.extend(docs)

        elif ext == ".docx":
            all_docs.extend(load_docx(path))

        elif ext == ".txt":
            all_docs.extend(load_txt(path))

        elif ext in [".pptx", ".ppt"]:
            all_docs.extend(load_pptx(path))


    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1200, chunk_overlap=200, separators=["\n\n", "\n"," ",""])

    splits = text_splitter.split_documents(all_docs)

    # Create vector store    
    os.makedirs(VECTOR_STORE_DIR,exist_ok=True)

    vectorstore = FAISS.from_documents(splits, embeddings)
    vectorstore.save_local(os.path.join(VECTOR_STORE_DIR, "faiss_index"))

    # Create BM25 retriever
    bm25_retriever = BM25Retriever.from_documents(splits)
    bm25_retriever.k = 5

     # Save BM25   
    with open(BM25_FILE_PATH, "wb") as f:
        pickle.dump(bm25_retriever, f)
    return (len(splits), len(file_paths))


def get_bm25_retriever():
    if not os.path.exists(BM25_FILE_PATH):
        return None

    try:
        with open(BM25_FILE_PATH,"rb") as f:
            return pickle.load(f)

    except Exception as e:
        print(f"BM25 load error: {e}")
        return None


     